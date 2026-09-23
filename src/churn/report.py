"""
Excel risk report: replaces the Streamlit dashboard (milestone 13's
original plan) with something a retention/ops team can just open - no
server, no Python environment, no localhost URL to share. Presents exactly
what the dashboard's "Look up a policy" and "Across the book" sections
showed - the business view, not the SHAP/threshold technical detail (that
stays in models/evaluation_report.json etc. for whoever needs it).

Sheet "Policy Detail": one row per scored policy - risk %, segment,
product/status, suggested action, and a plain-language "why flagged"
column - color-coded by segment, with a header autofilter so the reader
can sort/filter by risk themselves rather than scrolling a dropdown.

Sheet "Summary": total scored, how many are flagged at risk, and the same
segment/suggested-action breakdowns the dashboard charted, as native Excel
bar charts.

Deliberately does no ML work itself - reads tbl_policy_risk_score (already
computed by churn_predict.py, which as of this version only scores Active
policies - see filter_active_policies() in churn_predict.py), same "no
recomputing SHAP/clustering here" principle the dashboard followed.

Each run writes a dated file (risk_report_YYYYMMDD.xlsx) rather than
overwriting the same name, so a day's report doesn't clobber the previous
one.

Usage:
    python -m churn.report
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.engine import Engine

from churn.churn_db import get_engine
from churn.churn_trainer import MODEL_DIR
from churn.recommend import describe_factors


def default_report_path(today: date | None = None) -> Path:
    """models/risk_report_YYYYMMDD.xlsx for the given (or current) date - one file per day, not overwritten in place."""
    today = today or date.today()
    return MODEL_DIR / f"risk_report_{today:%Y%m%d}.xlsx"

LATEST_RISK_SCORES_QUERY = """
    SELECT r.policy_num, r.scored_at, r.churn_probability, r.segment_label,
           r.top_factors, r.recommended_action,
           p.insurance_product, p.payment_frequency, p.policy_status
    FROM tbl_policy_risk_score r
    JOIN tbl_policies p ON p.policy_num = r.policy_num
    WHERE r.scored_at = (SELECT MAX(scored_at) FROM tbl_policy_risk_score)
    ORDER BY r.churn_probability DESC
"""

# light fill colors, one per segment - same "High/Medium/Low risk" labels segmentation.py assigns
SEGMENT_FILL = {
    "High risk": PatternFill("solid", fgColor="FFC7CE"),
    "Medium risk": PatternFill("solid", fgColor="FFEB9C"),
    "Low risk": PatternFill("solid", fgColor="C6EFCE"),
}

DETAIL_COLUMN_RENAME = {
    "policy_num": "Policy",
    "churn_probability": "Churn Risk",
    "segment_label": "Segment",
    "insurance_product": "Product",
    "payment_frequency": "Payment Frequency",
    "policy_status": "Status",
    "recommended_action": "Suggested Action",
    "why_flagged": "Why Flagged",
}
DETAIL_COLUMN_WIDTHS = {
    "Policy": 14, "Churn Risk": 11, "Segment": 12, "Product": 14,
    "Payment Frequency": 16, "Status": 11, "Suggested Action": 55, "Why Flagged": 45,
}


def load_latest_risk_scores(engine: Engine | None = None) -> pd.DataFrame:
    """Pull the most recent scoring batch, joined with policy context - all scored policies, every segment."""
    return pd.read_sql(LATEST_RISK_SCORES_QUERY, engine or get_engine())


def build_policy_detail_table(scores: pd.DataFrame) -> pd.DataFrame:
    """One row per policy, business-facing columns only - no raw SHAP values or feature names."""
    detail = scores.copy()
    detail["why_flagged"] = detail["top_factors"].apply(
        lambda factors: "; ".join(describe_factors(factors)) or "-"
    )
    ordered_columns = list(DETAIL_COLUMN_RENAME)
    return detail[ordered_columns].rename(columns=DETAIL_COLUMN_RENAME)


def build_summary_tables(scores: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Segment counts, and suggested-action counts for the at-risk subset - the "across the book" numbers."""
    segment_counts = (
        scores["segment_label"].value_counts().rename_axis("Segment").reset_index(name="Policies")
    )

    at_risk = scores[scores["segment_label"] != "Low risk"]
    # each action string is "short headline - longer explanation" (see recommend.py) -
    # the headline alone makes a clean chart category label
    action_counts = (
        at_risk["recommended_action"].str.split(" - ").str[0]
        .value_counts().rename_axis("Suggested Action").reset_index(name="Policies")
    )
    return segment_counts, action_counts


def write_policy_detail_sheet(workbook: Workbook, detail: pd.DataFrame) -> Worksheet:
    sheet = workbook.active
    sheet.title = "Policy Detail"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="31333F")
    sheet.append(list(detail.columns))
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center")

    for _, row in detail.iterrows():
        sheet.append(list(row))

    segment_col = list(detail.columns).index("Segment") + 1
    risk_col = list(detail.columns).index("Churn Risk") + 1
    for row_idx in range(2, sheet.max_row + 1):
        segment = sheet.cell(row=row_idx, column=segment_col).value
        fill = SEGMENT_FILL.get(segment)
        if fill:
            sheet.cell(row=row_idx, column=segment_col).fill = fill
        sheet.cell(row=row_idx, column=risk_col).number_format = "0%"

    for i, column_name in enumerate(detail.columns, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = DETAIL_COLUMN_WIDTHS.get(column_name, 16)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    return sheet


def _add_bar_chart(sheet: Worksheet, title: str, header_row: int, row_count: int, anchor_row: int) -> None:
    chart = BarChart()
    chart.title = title
    chart.y_axis.title = "Policies"
    data = Reference(sheet, min_col=2, min_row=header_row, max_row=header_row + row_count)
    categories = Reference(sheet, min_col=1, min_row=header_row + 1, max_row=header_row + row_count)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.width, chart.height = 16, 8
    sheet.add_chart(chart, f"D{anchor_row}")


def write_summary_sheet(
    workbook: Workbook, scores: pd.DataFrame, segment_counts: pd.DataFrame, action_counts: pd.DataFrame
) -> Worksheet:
    sheet = workbook.create_sheet("Summary")
    bold = Font(bold=True)

    at_risk_count = len(scores[scores["segment_label"] != "Low risk"])
    sheet.append(["Policies scored", len(scores)])
    sheet.append(["Flagged at risk", f"{at_risk_count} ({at_risk_count / len(scores):.0%})"])
    sheet["A1"].font = bold
    sheet["A2"].font = bold
    sheet.append([])

    segment_header_row = sheet.max_row + 1
    sheet.append(list(segment_counts.columns))
    for cell in sheet[sheet.max_row]:
        cell.font = bold
    for _, row in segment_counts.iterrows():
        sheet.append(list(row))
    _add_bar_chart(sheet, "Policies by risk segment", segment_header_row, len(segment_counts), anchor_row=segment_header_row)

    sheet.append([])

    action_header_row = sheet.max_row + 1
    sheet.append(list(action_counts.columns))
    for cell in sheet[sheet.max_row]:
        cell.font = bold
    for _, row in action_counts.iterrows():
        sheet.append(list(row))
    _add_bar_chart(
        sheet, "Suggested actions for at-risk policies", action_header_row, len(action_counts),
        anchor_row=segment_header_row + 17,
    )

    sheet.column_dimensions["A"].width = 45
    sheet.column_dimensions["B"].width = 14
    return sheet


def build_report(scores: pd.DataFrame, path: Path | None = None) -> Path:
    """Build both sheets from an already-loaded scores table and save the workbook to path (default: today's dated filename)."""
    if scores.empty:
        raise ValueError("No scored policies to report on - run `python -m churn.churn_predict` first.")
    path = path or default_report_path()

    detail = build_policy_detail_table(scores)
    segment_counts, action_counts = build_summary_tables(scores)

    workbook = Workbook()
    write_policy_detail_sheet(workbook, detail)
    write_summary_sheet(workbook, scores, segment_counts, action_counts)

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


def run_report(engine: Engine | None = None, path: Path | None = None) -> Path:
    """Full pipeline: load the latest scoring batch from postgres, build the workbook, save it."""
    scores = load_latest_risk_scores(engine)
    return build_report(scores, path)


if __name__ == "__main__":
    saved_path = run_report()
    print(f"wrote risk report to {saved_path}")
