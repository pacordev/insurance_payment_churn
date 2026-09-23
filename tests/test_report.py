"""
Tests for report.py. build_policy_detail_table()/build_summary_tables() get
pure DataFrame tests (no db). build_report() is tested end-to-end against a
synthetic scores table, writing to pytest's tmp_path and reading the
workbook back with openpyxl - never touching the real models/ directory.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from churn.report import build_policy_detail_table, build_report, build_summary_tables, default_report_path


def _synthetic_scores() -> pd.DataFrame:
    return pd.DataFrame([
        {
            "policy_num": "POL-HIGH-1", "churn_probability": 0.91, "segment_label": "High risk",
            "top_factors": [{"feature": "failed_invoice_ratio", "value": 0.4},
                             {"feature": "tenure_days", "value": 0.3}],
            "recommended_action": "Prompt a payment-method update - repeated failed payment attempts...",
            "insurance_product": "Auto", "payment_frequency": "Monthly", "policy_status": "Active",
        },
        {
            "policy_num": "POL-MED-1", "churn_probability": 0.55, "segment_label": "Medium risk",
            "top_factors": [{"feature": "missed_invoice_ratio", "value": 0.1}],
            "recommended_action": "Proactive outreach before the next invoice is due - this policy...",
            "insurance_product": "Home", "payment_frequency": "Quarterly", "policy_status": "Active",
        },
        {
            "policy_num": "POL-LOW-1", "churn_probability": 0.02, "segment_label": "Low risk",
            "top_factors": [],
            "recommended_action": "Monitor - no single actionable behavioral driver stands out...",
            "insurance_product": "Life", "payment_frequency": "Monthly", "policy_status": "Active",
        },
    ])


def test_build_policy_detail_table_has_expected_columns_and_row_count() -> None:
    scores = _synthetic_scores()

    detail = build_policy_detail_table(scores)

    assert list(detail.columns) == [
        "Policy", "Churn Risk", "Segment", "Product", "Payment Frequency",
        "Status", "Suggested Action", "Why Flagged",
    ]
    assert len(detail) == 3


def test_build_policy_detail_table_translates_top_factors_to_plain_language() -> None:
    scores = _synthetic_scores()

    detail = build_policy_detail_table(scores).set_index("Policy")

    assert "Repeated failed payment attempts" in detail.loc["POL-HIGH-1", "Why Flagged"]
    assert detail.loc["POL-LOW-1", "Why Flagged"] == "-"


def test_build_summary_tables_counts_segments_and_at_risk_actions() -> None:
    scores = _synthetic_scores()

    segment_counts, action_counts = build_summary_tables(scores)

    assert dict(zip(segment_counts["Segment"], segment_counts["Policies"])) == {
        "High risk": 1, "Medium risk": 1, "Low risk": 1,
    }
    # the Low risk policy is excluded from the at-risk action breakdown
    assert action_counts["Policies"].sum() == 2
    assert "Monitor" not in set(action_counts["Suggested Action"])


def test_build_report_raises_a_clear_error_on_no_scored_policies() -> None:
    with pytest.raises(ValueError, match="churn_predict"):
        build_report(pd.DataFrame())


def test_build_report_writes_a_workbook_with_both_sheets(tmp_path: Path) -> None:
    scores = _synthetic_scores()

    path = build_report(scores, path=tmp_path / "risk_report.xlsx")

    assert path.exists()
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Policy Detail", "Summary"]

    detail_sheet = workbook["Policy Detail"]
    assert [cell.value for cell in detail_sheet[1]] == [
        "Policy", "Churn Risk", "Segment", "Product", "Payment Frequency",
        "Status", "Suggested Action", "Why Flagged",
    ]
    assert detail_sheet.max_row == 4  # header + 3 policies
    assert detail_sheet.auto_filter.ref is not None

    summary_sheet = workbook["Summary"]
    assert summary_sheet["A1"].value == "Policies scored"
    assert summary_sheet["B1"].value == 3


def test_default_report_path_uses_yyyymmdd_naming() -> None:
    path = default_report_path(today=date(2026, 3, 5))

    assert path.name == "risk_report_20260305.xlsx"
